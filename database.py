import streamlit as st

st.set_page_config(
    page_title="MiMercadito",
    page_icon="🛒",
    layout="wide"
)

import io
import sqlite3
import pandas as pd
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image, ImageDraw, ImageFont


# ============================================================
# BASE DE DATOS
# ============================================================
DB = "mercadito.db"

def conectar():
    conn = sqlite3.connect(DB, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn

def inicializar():
    conn = conectar()
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS productos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL UNIQUE,
            unidad TEXT DEFAULT 'kg',
            precio_compra REAL DEFAULT 0,
            precio_venta REAL NOT NULL,
            stock REAL DEFAULT 0,
            stock_minimo REAL DEFAULT 5
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS ventas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT NOT NULL,
            total REAL NOT NULL,
            metodo_pago TEXT DEFAULT 'efectivo',
            cliente TEXT DEFAULT ''
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS venta_detalle (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            venta_id INTEGER NOT NULL,
            producto_id INTEGER NOT NULL,
            nombre TEXT,
            cantidad REAL,
            precio_unitario REAL,
            subtotal REAL
        )
    """)
    conn.commit()
    conn.close()

def agregar_producto(nombre, precio_venta, unidad="kg",
                     precio_compra=0, stock=0, stock_minimo=5):
    conn = conectar()
    try:
        conn.execute(
            "INSERT INTO productos (nombre, unidad, precio_compra, "
            "precio_venta, stock, stock_minimo) VALUES (?,?,?,?,?,?)",
            (nombre, unidad, precio_compra, precio_venta, stock, stock_minimo)
        )
        conn.commit()
        return True
    except sqlite3.IntegrityError:
        return False
    finally:
        conn.close()

def listar_productos():
    conn = conectar()
    filas = conn.execute("SELECT * FROM productos ORDER BY nombre").fetchall()
    conn.close()
    return [dict(f) for f in filas]

def buscar_producto(texto):
    conn = conectar()
    filas = conn.execute(
        "SELECT * FROM productos WHERE nombre LIKE ? ORDER BY nombre",
        (f"%{texto}%",)
    ).fetchall()
    conn.close()
    return [dict(f) for f in filas]

def actualizar_producto(pid, **campos):
    if not campos:
        return
    sets = ", ".join(f"{k} = ?" for k in campos)
    vals = list(campos.values()) + [pid]
    conn = conectar()
    conn.execute(f"UPDATE productos SET {sets} WHERE id = ?", vals)
    conn.commit()
    conn.close()

def eliminar_producto(pid):
    conn = conectar()
    conn.execute("DELETE FROM productos WHERE id = ?", (pid,))
    conn.commit()
    conn.close()

def registrar_venta(items, metodo_pago="efectivo", cliente=""):
    conn = conectar()
    c = conn.cursor()
    total = sum(i["cantidad"] * i["precio"] for i in items)
    c.execute(
        "INSERT INTO ventas (fecha, total, metodo_pago, cliente) "
        "VALUES (?,?,?,?)",
        (datetime.now().isoformat(), total, metodo_pago, cliente)
    )
    venta_id = c.lastrowid
    for it in items:
        c.execute(
            "INSERT INTO venta_detalle "
            "(venta_id, producto_id, nombre, cantidad, "
            "precio_unitario, subtotal) VALUES (?,?,?,?,?,?)",
            (venta_id, it["id"], it["nombre"], it["cantidad"],
             it["precio"], it["cantidad"] * it["precio"])
        )
        c.execute(
            "UPDATE productos SET stock = stock - ? WHERE id = ?",
            (it["cantidad"], it["id"])
        )
    conn.commit()
    conn.close()
    return total

def ventas_del_dia(fecha=None):
    if fecha is None:
        fecha = datetime.now().strftime("%Y-%m-%d")
    conn = conectar()
    filas = conn.execute(
        "SELECT * FROM ventas WHERE date(fecha) = ? ORDER BY fecha DESC",
        (fecha,)
    ).fetchall()
    conn.close()
    return [dict(f) for f in filas]

def ventas_rango(desde, hasta):
    conn = conectar()
    filas = conn.execute(
        "SELECT * FROM ventas WHERE date(fecha) BETWEEN ? AND ? "
        "ORDER BY fecha DESC",
        (desde, hasta)
    ).fetchall()
    conn.close()
    return [dict(f) for f in filas]

def detalle_venta(venta_id):
    conn = conectar()
    filas = conn.execute(
        "SELECT * FROM venta_detalle WHERE venta_id = ?", (venta_id,)
    ).fetchall()
    conn.close()
    return [dict(f) for f in filas]

def productos_mas_vendidos(limite=10):
    conn = conectar()
    filas = conn.execute("""
        SELECT nombre, SUM(cantidad) AS total_cant,
               SUM(subtotal) AS total_ingreso
        FROM venta_detalle
        GROUP BY nombre
        ORDER BY total_cant DESC
        LIMIT ?
    """, (limite,)).fetchall()
    conn.close()
    return [dict(f) for f in filas]

def resumen_metodos_pago(fecha=None):
    if fecha is None:
        fecha = datetime.now().strftime("%Y-%m-%d")
    conn = conectar()
    filas = conn.execute("""
        SELECT metodo_pago, COUNT(*) AS n, SUM(total) AS total
        FROM ventas WHERE date(fecha) = ?
        GROUP BY metodo_pago
    """, (fecha,)).fetchall()
    conn.close()
    return [dict(f) for f in filas]


# ============================================================
# REPORTES
# ============================================================
def _fuente(tam):
    try:
        return ImageFont.truetype("DejaVuSans.ttf", tam)
    except Exception:
        return ImageFont.load_default()

def generar_excel(ventas, detalles):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "REPORTE DE VENTAS"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"

    total = sum(v["total"] for v in ventas)
    ws["A4"] = "Total vendido"; ws["B4"] = round(total, 2)
    ws["A5"] = "N de ventas"; ws["B5"] = len(ventas)
    ws["A6"] = "Ticket promedio"
    ws["B6"] = round(total / len(ventas), 2) if ventas else 0

    for fila in ws["A4:A6"]:
        fila[0].font = Font(bold=True)

    ws2 = wb.create_sheet("Detalle")
    ws2.append(["#", "Fecha", "Producto", "Cantidad",
                "P. Unit", "Subtotal", "Pago"])
    for celda in ws2[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2E7D32")
        celda.alignment = Alignment(horizontal="center")

    for i, v in enumerate(ventas, 1):
        for d in detalles.get(v["id"], []):
            ws2.append([
                i, v["fecha"][:16], d["nombre"], d["cantidad"],
                d["precio_unitario"], d["subtotal"], v["metodo_pago"]
            ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generar_imagen(ventas, top_productos, metodos):
    ancho, alto = 800, 1000
    img = Image.new("RGB", (ancho, alto), "white")
    draw = ImageDraw.Draw(img)

    titulo = _fuente(32)
    subt = _fuente(24)
    texto = _fuente(20)

    total = sum(v["total"] for v in ventas)
    y = 30

    draw.text((30, y), "REPORTE DEL DIA", fill="black", font=titulo)
    y += 55
    draw.text((30, y), f"Fecha: {datetime.now().strftime('%d/%m/%Y')}",
              fill="gray", font=texto)
    y += 45
    draw.line([(30, y), (770, y)], fill="black", width=2)
    y += 25
    draw.text((30, y), f"Total vendido:  S/ {total:.2f}",
              fill="#2E7D32", font=titulo)
    y += 55
    draw.text((30, y), f"N ventas:  {len(ventas)}",
              fill="black", font=subt)
    y += 40
    prom = total / len(ventas) if ventas else 0
    draw.text((30, y), f"Ticket promedio:  S/ {prom:.2f}",
              fill="black", font=subt)
    y += 60
    draw.text((30, y), "TOP PRODUCTOS", fill="black", font=subt)
    y += 45
    for i, p in enumerate(top_productos[:5], 1):
        linea = (f"{i}. {p['nombre']}  -  {p['total_cant']} u  "
                 f"-  S/ {p['total_ingreso']:.2f}")
        draw.text((40, y), linea, fill="black", font=texto)
        y += 35
    y += 30
    draw.text((30, y), "METODOS DE PAGO", fill="black", font=subt)
    y += 45
    for m in metodos:
        linea = (f"{m['metodo_pago'].upper()}:  S/ {m['total']:.2f}  "
                 f"({m['n']} ventas)")
        draw.text((40, y), linea, fill="black", font=texto)
        y += 35
    y += 40
    draw.text((30, y), "Generado con MiMercadito",
              fill="gray", font=texto)

    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


# ============================================================
# APP
# ============================================================
inicializar()

if "carrito" not in st.session_state:
    st.session_state.carrito = []

st.title("🛒 MiMercadito")
st.caption(f"📅 {datetime.now().strftime('%d/%m/%Y %H:%M')}")

tab_venta, tab_productos, tab_reportes, tab_historial = st.tabs(
    ["🧾 Vender", "📦 Productos", "📊 Reportes", "📜 Historial"]
)

# ---------- VENDER ----------
with tab_venta:
    filtro = st.text_input("🔍 Buscar producto",
                           placeholder="Ej: papa, arroz, cebolla")
    productos = buscar_producto(filtro) if filtro else listar_productos()

    if not productos:
        st.warning("No hay productos. Ve a **Productos** para agregar.")
    else:
        st.markdown("### Productos disponibles")
        cols = st.columns(3)
        for i, p in enumerate(productos):
            with cols[i % 3]:
                with st.container(border=True):
                    st.markdown(f"**{p['nombre']}**")
                    st.caption(f"S/ {p['precio_venta']:.2f} / {p['unidad']}")
                    st.caption(f"Stock: {p['stock']:.1f} {p['unidad']}")

                    cantidad = st.number_input(
                        "Cantidad", min_value=0.1, value=1.0, step=0.5,
                        key=f"cant_{p['id']}", label_visibility="collapsed"
                    )
                    if st.button("➕ Agregar", key=f"add_{p['id']}"):
                        existe = False
                        for item in st.session_state.carrito:
                            if item["id"] == p["id"]:
                                item["cantidad"] += cantidad
                                existe = True
                                break
                        if not existe:
                            st.session_state.carrito.append({
                                "id": p["id"], "nombre": p["nombre"],
                                "cantidad": cantidad,
                                "precio": p["precio_venta"],
                                "unidad": p["unidad"],
                            })
                        st.rerun()

    st.divider()
    st.markdown("### 🛒 Carrito")

    if st.session_state.carrito:
        for idx, item in enumerate(st.session_state.carrito):
            col1, col2, col3, col4 = st.columns([4, 1, 1, 1])
            with col1:
                st.write(f"**{item['nombre']}** — "
                         f"{item['cantidad']} {item['unidad']} × "
                         f"S/ {item['precio']:.2f}")
            with col2:
                st.write(f"S/ {item['cantidad'] * item['precio']:.2f}")
            with col3:
                if st.button("➖", key=f"menos_{idx}"):
                    item["cantidad"] -= 0.5
                    if item["cantidad"] <= 0:
                        st.session_state.carrito.pop(idx)
                    st.rerun()
            with col4:
                if st.button("❌", key=f"del_{idx}"):
                    st.session_state.carrito.pop(idx)
                    st.rerun()

        total = sum(i["cantidad"] * i["precio"]
                    for i in st.session_state.carrito)
        st.metric("💰 Total a cobrar", f"S/ {total:.2f}")

        col1, col2, col3 = st.columns(3)
        with col1:
            metodo = st.selectbox("Método de pago",
                                  ["efectivo", "yape", "plin", "transferencia"])
        with col2:
            cliente = st.text_input("Cliente (opcional)", "")
        with col3:
            st.write("")
            st.write("")
            if st.button("✅ COBRAR", type="primary"):
                registrar_venta(st.session_state.carrito,
                                metodo_pago=metodo, cliente=cliente)
                st.success(f"✅ Venta registrada: S/ {total:.2f}")
                st.session_state.carrito = []
                st.rerun()

        if st.button("🗑️ Vaciar carrito"):
            st.session_state.carrito = []
            st.rerun()
    else:
        st.info("El carrito está vacío. Agrega productos arriba.")

# ---------- PRODUCTOS ----------
with tab_productos:
    st.markdown("### ➕ Nuevo producto")
    with st.form("nuevo_producto", clear_on_submit=True):
        c1, c2 = st.columns(2)
        with c1:
            nombre = st.text_input("Nombre")
            unidad = st.selectbox("Unidad", ["kg", "unidad", "litro", "docena"])
            precio_venta = st.number_input("Precio de venta (S/)",
                                           min_value=0.0, step=0.1)
        with c2:
            precio_compra = st.number_input("Precio de compra (S/)",
                                            min_value=0.0, step=0.1)
            stock = st.number_input("Stock inicial", min_value=0.0, step=1.0)
            stock_minimo = st.number_input("Stock mínimo",
                                           min_value=0.0, value=5.0)
        enviado = st.form_submit_button("💾 Guardar producto", type="primary")
        if enviado:
            if not nombre:
                st.error("El nombre es obligatorio")
            elif agregar_producto(nombre, precio_venta, unidad,
                                  precio_compra, stock, stock_minimo):
                st.success(f"✅ Producto '{nombre}' agregado")
                st.rerun()
            else:
                st.error("Ese producto ya existe")

    st.divider()
    st.markdown("### 📋 Lista de productos")
    productos = listar_productos()
    if productos:
        df = pd.DataFrame(productos)
        df = df[["id", "nombre", "unidad", "precio_compra",
                 "precio_venta", "stock", "stock_minimo"]]
        df.columns = ["ID", "Nombre", "Unidad", "P. Compra",
                      "P. Venta", "Stock", "Stock mín."]
        st.dataframe(df, use_container_width=True, hide_index=True)

        st.markdown("#### ✏️ Editar o eliminar")
        opciones = {f"{p['id']} - {p['nombre']}": p for p in productos}
        sel = st.selectbox("Selecciona un producto", list(opciones.keys()))
        if sel:
            p = opciones[sel]
            with st.form("editar"):
                c1, c2, c3 = st.columns(3)
                with c1:
                    nuevo_precio = st.number_input(
                        "Precio venta", value=float(p["precio_venta"]),
                        min_value=0.0, step=0.1)
                    nuevo_stock = st.number_input(
                        "Stock", value=float(p["stock"]),
                        min_value=0.0, step=1.0)
                with c2:
                    nuevo_pc = st.number_input(
                        "Precio compra", value=float(p["precio_compra"]),
                        min_value=0.0, step=0.1)
                    nuevo_sm = st.number_input(
                        "Stock mínimo", value=float(p["stock_minimo"]),
                        min_value=0.0, step=1.0)
                with c3:
                    unidades_validas = ["kg", "unidad", "litro", "docena"]
                    idx_u = (unidades_validas.index(p["unidad"])
                             if p["unidad"] in unidades_validas else 0)
                    nueva_unidad = st.selectbox("Unidad", unidades_validas,
                                                index=idx_u)
                col_a, col_b = st.columns(2)
                with col_a:
                    if st.form_submit_button("💾 Actualizar", type="primary"):
                        actualizar_producto(
                            p["id"], precio_venta=nuevo_precio,
                            precio_compra=nuevo_pc, stock=nuevo_stock,
                            stock_minimo=nuevo_sm, unidad=nueva_unidad)
                        st.success("Actualizado")
                        st.rerun()
                with col_b:
                    if st.form_submit_button("🗑️ Eliminar"):
                        eliminar_producto(p["id"])
                        st.warning("Eliminado")
                        st.rerun()
    else:
        st.info("Aún no hay productos registrados.")

# ---------- REPORTES ----------
with tab_reportes:
    st.markdown("### 📊 Reportes")
    opcion = st.radio("Rango",
                      ["Hoy", "Últimos 7 días", "Este mes", "Personalizado"],
                      horizontal=True)
    hoy = date.today()
    if opcion == "Hoy":
        desde = hasta = hoy
    elif opcion == "Últimos 7 días":
        desde = date.fromordinal(hoy.toordinal() - 6)
        hasta = hoy
    elif opcion == "Este mes":
        desde = hoy.replace(day=1)
        hasta = hoy
    else:
        c1, c2 = st.columns(2)
        with c1:
            desde = st.date_input("Desde", hoy)
        with c2:
            hasta = st.date_input("Hasta", hoy)

    ventas = ventas_rango(str(desde), str(hasta))
    detalles = {v["id"]: detalle_venta(v["id"]) for v in ventas}
    top = productos_mas_vendidos()
    metodos = resumen_metodos_pago(str(hasta))

    total = sum(v["total"] for v in ventas)
    n = len(ventas)
    prom = total / n if n else 0

    c1, c2, c3 = st.columns(3)
    c1.metric("💰 Total vendido", f"S/ {total:.2f}")
    c2.metric("🧾 N ventas", n)
    c3.metric("📊 Ticket promedio", f"S/ {prom:.2f}")

    st.markdown("#### 🏆 Top productos")
    if top:
        df_top = pd.DataFrame(top)
        df_top.columns = ["Producto", "Cantidad", "Ingreso (S/)"]
        st.dataframe(df_top, use_container_width=True, hide_index=True)
    else:
        st.info("Sin ventas aún")

    st.markdown("#### 💳 Métodos de pago")
    if metodos:
        df_m = pd.DataFrame(metodos)
        df_m.columns = ["Método", "N ventas", "Total (S/)"]
        st.dataframe(df_m, use_container_width=True, hide_index=True)

    st.divider()
    st.markdown("#### 📤 Exportar")
    c1, c2 = st.columns(2)
    with c1:
        excel = generar_excel(ventas, detalles)
        st.download_button(
            "📊 Descargar Excel", data=excel,
            file_name=f"reporte_{desde}_{hasta}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
    with c2:
        imagen = generar_imagen(ventas, top, metodos)
        st.download_button(
            "🖼️ Descargar Imagen", data=imagen,
            file_name=f"reporte_{desde}_{hasta}.png",
            mime="image/png", use_container_width=True
        )

# ---------- HISTORIAL ----------
with tab_historial:
    st.markdown("### 📜 Historial de ventas")
    ventas = ventas_del_dia()
    if ventas:
        for v in ventas:
            with st.expander(
                f"Venta #{v['id']} — S/ {v['total']:.2f} — "
                f"{v['metodo_pago']} — {v['fecha'][:16]}"
            ):
                detalles_v = detalle_venta(v["id"])
                if detalles_v:
                    df = pd.DataFrame(detalles_v)
                    df = df[["nombre", "cantidad",
                             "precio_unitario", "subtotal"]]
                    df.columns = ["Producto", "Cantidad", "P. Unit", "Subtotal"]
                    st.dataframe(df, use_container_width=True,
                                 hide_index=True)
                if v.get("cliente"):
                    st.caption(f"👤 Cliente: {v['cliente']}")
    else:
        st.info("No hay ventas registradas hoy.")
